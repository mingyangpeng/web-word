#!/usr/bin/env python3
"""导出「域系初判表 v0.2」——义位筛选规则表（352 保留义位）按域系初判表头输出。

数据链：
  物理位移义位表（358 候选，机械正则）
    → 义位筛选规则表（R1 剔除 6 / R2 明确保留 5 / R3 边界复核 3 / R0 常规 344）
    → 本表：352 保留义位 × 域系初判表 02 的表头与词目初判归属

口径：仍为【词目层初判】预览（义位图式暂取词目初判），最终域—系体系以
义位层人工标注后重聚类为准。与 v0.1 的差异：剔除 6 条 R1 抽象义位后的净版。

产出：docs/域系初判表v0.2.xlsx
  01 域—系总览（含域小计、代表例句）
  02 逐义位明细（表头与 v0.1 完全一致 + 规则列）
  03 口径说明
"""

import os
import sys
from datetime import date
from pathlib import Path
from collections import defaultdict

XLSX_SKILL_DIR = "/Users/jhon/.zcode/cli/plugins/cache/zcode-plugins-official/spreadsheets/0.1.7/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)
from base import (  # noqa: E402
    FONT_NAME, FORMATS, PRIMARY_LIGHT, setup_sheet, style_header_row,
    style_data_row, style_total_row, font_caption, font_subheader,
    auto_fit_columns, auto_fit_row_heights,
)
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "域系初判表v0.2.xlsx"

# ---- 数据 ----
wb_rule = load_workbook(ROOT / "docs" / "义位筛选规则表.xlsx", read_only=True)
kept = []
for r in wb_rule["03 保留义位全量"].iter_rows(min_row=5, values_only=True):
    if r[1] is None or r[1] == "合计":
        break
    # 03 表列序（B 起）：SenseID, 词目, 语法码, 英文释义, 例句, 规则, 备注
    kept.append({"sid": str(r[1]), "verb": str(r[2]), "grammar": str(r[3] or ""),
                 "defn": str(r[4] or ""), "ex": str(r[5] or ""), "rule": str(r[6] or "")})
wb_rule.close()

wb_dom = load_workbook(ROOT / "docs" / "域系初判表.xlsx", read_only=True)
assign = {}
for r in wb_dom["02 逐义位明细"].iter_rows(min_row=5, values_only=True):
    if r[1] is None or r[1] == "合计":
        continue
    assign[str(r[1])] = {"schema": r[5], "domain": r[6], "family": r[7]}
wb_dom.close()

for k in kept:
    a = assign.get(k["sid"], {"schema": "?", "domain": "?", "family": "?"})
    k.update(a)

by_dom = defaultdict(lambda: defaultdict(list))
for k in kept:
    by_dom[k["domain"]][k["family"]].append(k)

# ---- 工作簿 ----
wb = Workbook()

# 01 域—系总览
ws = wb.active
ws.title = "01 域系总览"
H1 = ["认知域（初判）", "语义系（初判）", "义位数", "代表义位与例句"]
setup_sheet(ws, title="域—系初判表 v0.2（352 保留义位 · R1 剔除后净版 · 词目层初判）", last_col=len(H1) + 1)
for col, h in enumerate(H1, 2):
    ws.cell(row=4, column=col, value=h)
style_header_row(ws, row_num=4, col_start=2, col_end=len(H1) + 1)
rn, i = 5, 0
for dom in ["方式位移域", "路径位移域", "互动位移域"]:
    subs = by_dom.get(dom, {})
    for sub, items in sorted(subs.items(), key=lambda x: -len(x[1])):
        exs = []
        for it in items:
            if it["ex"] and it["ex"] != "None" and len(exs) < 2:
                exs.append(f"{it['sid']}: {it['ex'].split(chr(10))[0]}")
        ws.cell(row=rn, column=2, value=dom)
        ws.cell(row=rn, column=3, value=sub)
        ws.cell(row=rn, column=4, value=len(items))
        ws.cell(row=rn, column=5, value="\n".join(exs))
        style_data_row(ws, row_num=rn, col_start=2, col_end=len(H1) + 1, row_index=i)
        ws.cell(row=rn, column=4).number_format = FORMATS["integer"]
        rn += 1
        i += 1
    dom_total = sum(len(v) for v in subs.values())
    ws.cell(row=rn, column=2, value=f"{dom} 小计")
    ws.cell(row=rn, column=4, value=dom_total)
    style_total_row(ws, row_num=rn, col_start=2, col_end=len(H1) + 1)
    ws.cell(row=rn, column=4).number_format = FORMATS["integer"]
    rn += 1
ws.cell(row=rn, column=2, value="合计")
ws.cell(row=rn, column=4, value=len(kept))
style_total_row(ws, row_num=rn, col_start=2, col_end=len(H1) + 1)
ws.cell(row=rn, column=4).number_format = FORMATS["integer"]
ws.cell(row=rn + 2, column=2,
        value="注：v0.2 = v0.1（358 候选）剔除 R1 抽象义位 6 条后的净版；仍为词目层初判预览，"
              "最终体系以义位层人工标注后重聚类为准。").font = font_caption()
ws.column_dimensions["E"].width = 80
auto_fit_columns(ws, min_width=8, max_width=30, header_row=4, data_start_row=5)
auto_fit_row_heights(ws, header_row=4, data_start_row=5)

# 02 逐义位明细（表头与 v0.1 一致 + 规则列）
ws2 = wb.create_sheet("02 逐义位明细")
H2 = ["SenseID", "词目", "英文释义", "例句（首条）", "词目层初判图式", "域（初判）", "系（初判）",
      "义位层标注（待人工）", "筛选规则"]
setup_sheet(ws2, title="352 保留义位逐条明细（含域—系初判归属与筛选规则标签）", last_col=len(H2) + 1)
for col, h in enumerate(H2, 2):
    ws2.cell(row=4, column=col, value=h)
style_header_row(ws2, row_num=4, col_start=2, col_end=len(H2) + 1)
order = {"方式位移域": 0, "路径位移域": 1, "互动位移域": 2}
FILL_R2 = PatternFill("solid", fgColor="E8F5E9")
FILL_R3 = PatternFill("solid", fgColor="FEF9E7")
for i, k in enumerate(sorted(kept, key=lambda x: (order.get(x["domain"], 9), x["family"], x["verb"]))):
    rn = 5 + i
    first_ex = k["ex"].split("\n")[0] if k["ex"] and k["ex"] != "None" else ""
    vals = [k["sid"], k["verb"], k["defn"], first_ex, k["schema"], k["domain"], k["family"],
            None, k["rule"] or "R0"]
    for col, v in enumerate(vals, 2):
        ws2.cell(row=rn, column=col, value=v)
    style_data_row(ws2, row_num=rn, col_start=2, col_end=len(H2) + 1, row_index=i)
    ws2.cell(row=rn, column=9).fill = PatternFill("solid", fgColor=PRIMARY_LIGHT)
    if k["rule"].startswith("R2"):
        ws2.cell(row=rn, column=10).fill = FILL_R2
    elif k["rule"].startswith("R3"):
        ws2.cell(row=rn, column=10).fill = FILL_R3
total = 5 + len(kept)
ws2.cell(row=total, column=2, value="合计")
ws2.cell(row=total, column=3, value=f"{len(kept)} 义位")
style_total_row(ws2, row_num=total, col_start=2, col_end=len(H2) + 1)
ws2.freeze_panes = "D5"
ws2.auto_filter.ref = f"B4:J{total - 1}"
auto_fit_columns(ws2, min_width=8, max_width=56, header_row=4, data_start_row=5)
auto_fit_row_heights(ws2, header_row=4, data_start_row=5)

# 03 口径说明
ws3 = wb.create_sheet("03 口径说明")
setup_sheet(ws3, title="域系初判表 v0.2 · 口径说明", last_col=3)
ws3.column_dimensions["B"].width = 24
ws3.column_dimensions["C"].width = 100
rn = 4
for k, v in [
    ("本表性质", "词目层初判预览（v0.2）：义位图式暂取词目初判，非义位层标注结果，"
             "不能直接作为论文 4.2.2 的最终体系。"),
    ("与 v0.1 差异", "v0.1 基于 358 条物理位移候选；v0.2 在其上执行义位筛选规则"
                "（R1 剔除 6 条抽象义位：run-S35、rush-S6/S7、race-S1/S2/S5），"
                "保留 352 条（R0 常规 344 + R2 明确保留 5 + R3 边界复核 3）。"),
    ("规则标签列", "02 表新增「筛选规则」列：R2 绿底＝用户明确保留的隐喻性空间位移；"
              "R3 琥珀底＝边界案例（保留待双重核验）；R0＝常规物理位移。"),
    ("域系分布变化", "方式位移域 270→264（奔跑行进系 56→50，剔除集中在 run/race/rush 的抽象义）；"
                "路径位移域 62、互动位移域 26 不变。"),
    ("下一步", "352 条进 LLM 初筛（等 ZHIPUAI_API_KEY）→ 人工逐条复核 → Y 类义位按实际图式"
           "标签重新聚类 → 升级 v1.0 正式体系表（重跑本脚本）。"),
    ("生成信息", f"生成日期 {date.today()}；脚本 scripts/export_domain_family_v02.py 可重跑。"),
]:
    ws3.cell(row=rn, column=2, value=k)
    ws3.cell(row=rn, column=3, value=v)
    style_data_row(ws3, row_num=rn, col_start=2, col_end=3, row_index=rn % 2)
    rn += 1
auto_fit_row_heights(ws3, header_row=4, data_start_row=4)

wb.properties.creator = "Z.ai"
wb.save(OUT)
print(f"saved: {OUT}")
print(f"total={len(kept)}; 方式{sum(len(v) for s in by_dom['方式位移域'].values() for v in [s])}"
      f"/路径{sum(len(v) for s in by_dom['路径位移域'].values() for v in [s])}"
      f"/互动{sum(len(v) for s in by_dom['互动位移域'].values() for v in [s])}")
