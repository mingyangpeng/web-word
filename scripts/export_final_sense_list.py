#!/usr/bin/env python3
"""导出【表3 最终研究义位表】——三表工作流的结果层（进入 MOVECAT 核心分析的义位集合）。

来源：
  - 应用库 oald_senses：human_reviewed=1 且 prototype_level ∈ {Y, 拓展} 的义位（平台人工标注结果）
  - output/02_义位标注/*_义项.xlsx 人工列：判定理由 / 备注（平台暂无理由字段时的理由层，SenseID 关联）
  - corpus_word_stats：enTenTen21 词目频次

产出 docs/表3_最终研究义位表.xlsx：
  「核心研究义位(Y)」完整五维标注 + 理由，进入三域七系范畴聚类与释义模板构建；
  「隐喻拓展义位(拓展)」简化字段，用于原型辐射网络与英汉隐喻对比；
  「状态与说明」行数、与表2 的口径核对、缺失项清单。

标注完成（或进行中）后重跑：python3 scripts/export_final_sense_list.py
"""

import sys
import os
from datetime import date
from pathlib import Path

import yaml
from openpyxl import load_workbook

XLSX_SKILL_DIR = "/Users/jhon/.zcode/cli/plugins/cache/zcode-plugins-official/spreadsheets/0.1.7/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)
from base import (  # noqa: E402
    FONT_NAME, PRIMARY_LIGHT, setup_sheet, style_header_row, style_data_row,
    style_total_row, font_caption, font_subheader, auto_fit_columns,
    auto_fit_row_heights,
)
from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "表3_最终研究义位表.xlsx"
SENSE_DIR = ROOT / "output" / "02_义位标注"

# ---- 表2 人工列（理由层，SenseID 关联）
reasons = {}
conflicts = []
for f in sorted(SENSE_DIR.glob("*_义项.xlsx")):
    wb = load_workbook(f, read_only=True)
    ws = wb["义位标注"]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0]:
            continue
        sid, m, n, o = r[0], r[12], r[13], r[14]  # A SenseID / M 准入 / N 理由 / O 备注
        if m or n or o:
            reasons[sid] = (m or "", n or "", o or "")
    wb.close()

# ---- 平台标注结果（DB）
cfg = yaml.safe_load((ROOT / "config" / "config.yaml").read_text())["database"]["mysql"]
host = cfg["host"] if cfg["host"] != "mysql" else "127.0.0.1"
import pymysql  # noqa: E402

conn = pymysql.connect(host=host, port=cfg["port"], user=cfg["user"],
                       password=cfg["password"], database=cfg["database"],
                       charset="utf8mb4", cursorclass=pymysql.cursors.Cursor)
with conn.cursor() as cur:
    cur.execute(
        "SELECT headword, sense_index, definition, grammar, is_idiom, "
        "prototype_level, image_schema, alignment, cn_syntax, translation_strategy "
        "FROM oald_senses WHERE human_reviewed = 1 AND prototype_level IN ('Y','拓展')")
    annotated = cur.fetchall()
    cur.execute("SELECT verb, freq FROM corpus_word_stats")
    freq = dict(cur.fetchall())
conn.close()

def sense_id(headword, idx):
    return f"{headword}-S{int(idx)}"

y_rows, e_rows = [], []
for hw, idx, definition, grammar, is_idiom, level, schema, align, cn_syn, strategy in annotated:
    sid = sense_id(hw, idx)
    reason = reasons.get(sid, ("", "", ""))
    if reason[0] and reason[0] != level:
        conflicts.append((sid, reason[0], level))
    row = [sid, hw, f"S{int(idx)}", grammar or "", definition or "",
           level, schema or "", align or "", cn_syn or "", strategy or "",
           reason[1], reason[2], freq.get(hw, 0)]
    (y_rows if level == "Y" else e_rows).append(row)

def build_sheet(ws, title, headers, data):
    LAST = len(headers) + 1
    setup_sheet(ws, title=title, last_col=LAST)
    for col, h in enumerate(headers, 2):
        ws.cell(row=4, column=col, value=h)
    style_header_row(ws, row_num=4, col_start=2, col_end=LAST)
    for i, r in enumerate(data):
        rn = 5 + i
        for col, v in enumerate(r, 2):
            ws.cell(row=rn, column=col, value=v)
        style_data_row(ws, row_num=rn, col_start=2, col_end=LAST, row_index=i)
    total = 5 + len(data)
    ws.cell(row=total, column=2, value="合计")
    ws.cell(row=total, column=3, value=f"{len(data)} 义位")
    style_total_row(ws, row_num=total, col_start=2, col_end=LAST)
    ws.freeze_panes = "D5"
    auto_fit_columns(ws, min_width=8, max_width=40, header_row=4, data_start_row=5)
    auto_fit_row_heights(ws, header_row=4, data_start_row=5)
    return total

wb = Workbook()

H_FULL = ["本研究义位ID", "词形", "OALD义项号", "及物性", "英文释义", "准入标签",
          "意象图式", "英汉对齐等级", "汉语句法特征", "译义转化策略",
          "判定理由", "备注", "词目频次(enTenTen21)"]
ws = wb.active
ws.title = "核心研究义位(Y)"
t1 = build_sheet(ws, "表3 · 核心研究义位（准入标签 = Y，进入范畴聚类与释义模板构建）",
                 H_FULL, y_rows)

ws2 = wb.create_sheet("隐喻拓展义位(拓展)")
H_EXT = ["本研究义位ID", "词形", "OALD义项号", "及物性", "英文释义", "准入标签",
         "意象图式（简标）", "判定理由", "备注", "词目频次(enTenTen21)"]
ext_headers_idx = [0, 1, 2, 3, 4, 5, 6, 10, 11, 12]
e_data = [[r[i] for i in ext_headers_idx] for r in e_rows]
build_sheet(ws2, "表3 · 隐喻拓展义位（准入标签 = 拓展，仅用于语义辐射网络与英汉隐喻对比）",
            H_EXT, e_data)

# ---- 状态与说明
ws3 = wb.create_sheet("状态与说明")
setup_sheet(ws3, title="表3 状态与说明", last_col=3)
ws3.column_dimensions["B"].width = 24
ws3.column_dimensions["C"].width = 104
rn = 4

def sec(title):
    global rn
    ws3.cell(row=rn, column=2, value=title).font = font_subheader()
    for col in (2, 3):
        ws3.cell(row=rn, column=col).fill = PatternFill("solid", fgColor=PRIMARY_LIGHT)
    ws3.row_dimensions[rn].height = 24
    rn += 1

def kv(k, v, warn=False):
    global rn
    ws3.cell(row=rn, column=2, value=k)
    ws3.cell(row=rn, column=3, value=v)
    style_data_row(ws3, row_num=rn, col_start=2, col_end=3, row_index=rn % 2)
    if warn:
        ws3.cell(row=rn, column=3).font = Font(name=FONT_NAME, size=11, color="C0392B")
    rn += 1

sec("一、生成状态")
kv("生成日期", str(date.today()))
kv("数据来源", "平台 oald_senses（human_reviewed=1）＋ 02_义位标注 人工列（理由层）＋ corpus_word_stats（频次）")
kv("核心研究义位（Y）", f"{len(y_rows)} 条", warn=(len(y_rows) == 0))
kv("隐喻拓展义位（拓展）", f"{len(e_rows)} 条", warn=(len(e_rows) == 0))
kv("表2 已写人工列的义位", f"{len(reasons)} 条（表2 xlsx 侧，用于理由关联与口径核对）")
kv("准入口径冲突（表2 vs 平台）", f"{len(conflicts)} 条" + (f"：{conflicts[:5]}" if conflicts else ""),
   warn=bool(conflicts))
rn += 1

sec("二、口径说明")
kv("工作流位置", "表1（原始总表，只读）→ 表2（02_义位标注 人工列，逐条判定写理由）→ 表3（本文件，"
              "从平台人工标注结果提取）。本表只收录人工复核完成的义位（human_reviewed=1），"
              "平台未复核的义位不会进入。")
kv("Y 与拓展的分工", "Y＝核心原型物理位移义：完整五维标注，进入三域七系范畴聚类、多维释义模板构建与核心统计；"
                 "拓展＝隐喻引申义：单独成页，仅用于原型辐射网络描写与英汉隐喻图式对比，不进核心统计。")
kv("已知边界", "4 个补充词（scutter / kayak / oar / dodder）的义位不在平台库（平台载 oald_senses 665 条，"
            "词表合计 669 条）；如最终筛选涉及补充词，需先在平台补录或从表2 xlsx 人工导出。"
            "习语义项按编码手册一般归 N 或拓展，剔除与否在词表义项清单「问题清单」页确认。")
kv("重跑方式", "python3 scripts/export_final_sense_list.py（标注进行中可随时重跑刷新本表）。")
auto_fit_row_heights(ws3, header_row=4, data_start_row=4)

wb.properties.creator = "Z.ai"
wb.save(OUT)
print(f"saved: {OUT}")
print(f"Y={len(y_rows)}, ext={len(e_rows)}, reasons_in_xlsx={len(reasons)}, conflicts={len(conflicts)}")
