#!/usr/bin/env python3
"""生成 171 动词的意象图式提取结果数据集（AI 初筛层）。

对每个动词，按四类意象图式的证据槽位映射（《03 意象图式提取说明》2.1），
从 wsketch_full 原始 JSON 中聚合：
  - 覆盖槽位数 / 覆盖率
  - 证据搭配词数 N（去重词型，跨槽位取最大 logDice）
  - Mean_logDice
  - 重心得分 S = Mean_logDice × log2(N+1)（与 PS 同族公式）
  - 每图式 Top10 证据搭配词（含所在槽位）

产出（output/06_图式提取结果/）：
  - 图式提取结果汇总.xlsx（宽表矩阵 + 长表明细 + 说明）
  - schema_evidence/<动词>.json（逐词机器可读结果）

⚠ 定位：本结果为**未做语义归类的原始版**（AI 初筛），仅按槽位映射聚合；
  按《义项判定标准修改稿》，语义归类（剔除抽象/隐喻搭配词）后须重算，
  最终图式标签由研究者人工判定。公式与局限见方法说明文档。
"""
import json
import math
import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(BASE, "data", "corpus", "wsketch_full")
VERBS_FILE = os.path.join(BASE, "data", "corpus", "verbs_171.json")
OUT_DIR = os.path.join(BASE, "output", "06_图式提取结果")
JSON_DIR = os.path.join(OUT_DIR, "schema_evidence")

SCHEMA_SLOTS = {
    "源点脱离图式": [
        'particles after "%w"', 'objects of "%w away"', 'objects of "%w off"',
        'objects of "%w out"', 'prepositional phrases',
    ],
    "终点抵达图式": [
        'objects of "%w"', 'objects of "%w up"', 'objects of "%w down"',
        'objects of "%w in"', 'objects of "%w into"', 'prepositional phrases',
        'particles after "%w" with object',
    ],
    "方式伴随图式": [
        'modifiers of "%w"', 'adjectives after "%w"', 'usage patterns',
        'subjects of "%w"', '"%w" and/or ...', 'complements of "%w"',
    ],
    "互动位移图式": [
        'objects of "%w"', 'pronominal objects of "%w"', 'pronominal subjects of "%w"',
        'subjects of "%w"', 'prepositional phrases', 'particles after "%w" with object',
    ],
}

HEADER_FILL = PatternFill("solid", fgColor="4472C4")
HEADER_FONT = Font(bold=True, color="FFFFFF")
CAND_FILL = PatternFill("solid", fgColor="FFF2CC")
CENTER = Alignment(horizontal="center", vertical="center")
WRAP = Alignment(vertical="top", wrap_text=True)


def slot_entries(rel):
    """返回槽位的搭配条目 [(词形, logDice), ...]；嵌套槽位（PP 子槽）直接取子槽名。"""
    words = rel.get("Words", [])
    if words and "word" not in words[0] and "Words" in words[0]:
        return [(sub.get("name"), sub.get("score") or 0) for sub in words]
    return [(w.get("word"), w.get("score") or 0) for w in words]


def extract_verb(verb):
    j = json.load(open(os.path.join(SRC, verb + ".json"), encoding="utf-8"))
    rel_map = {r["name"]: r for r in j.get("Gramrels", [])}
    schemas = {}
    for schema, slots in SCHEMA_SLOTS.items():
        collocates = {}  # 词形 -> (max logDice, 槽位)
        slot_details = []
        for slot in slots:
            rel = rel_map.get(slot)
            if not rel:
                continue
            entries = slot_entries(rel)
            for w, sc in entries:
                if w is None:
                    continue
                if w not in collocates or sc > collocates[w][0]:
                    collocates[w] = (sc, slot)
            slot_details.append({
                "slot": slot, "slot_freq": rel.get("count"), "slot_logdice": rel.get("score"),
                "n_collocates": len(entries),
                "top5": [{"word": w, "logdice": sc} for w, sc in sorted(entries, key=lambda x: -x[1])[:5]],
            })
        n = len(collocates)
        scores = [v[0] for v in collocates.values()]
        mean_ld = sum(scores) / n if n else 0.0
        s_score = mean_ld * math.log2(n + 1) if n else 0.0
        top = sorted(collocates.items(), key=lambda x: -x[1][0])[:10]
        schemas[schema] = {
            "slots_present": len(slot_details), "slots_defined": len(slots),
            "coverage": round(len(slot_details) / len(slots), 3),
            "n_types": n, "mean_logdice": round(mean_ld, 3),
            "prominence_score": round(s_score, 3),
            "top_evidence": [{"word": w, "logdice": v[0], "slot": v[1]} for w, v in top],
            "slot_details": slot_details,
        }
    ranked = sorted(schemas.items(), key=lambda kv: -kv[1]["prominence_score"])
    return {
        "verb": verb, "freq": j.get("freq"), "lpos": j.get("lpos"),
        "corpus": "enTenTen21", "retrieved": "2026-09-30",
        "formulas": {
            "prominence_score": "S = Mean_logDice × log2(N+1)，N 为该图式证据槽位去重搭配词型数（未语义归类）",
            "note": "AI 初筛原始版：未做语义归类，结果仅作候选参考，最终图式标签须人工判定（义项判定标准修改稿 标准2）",
        },
        "candidate_schema": ranked[0][0] if ranked and ranked[0][1]["n_types"] > 0 else None,
        "ranking": [k for k, v in ranked],
        "schemas": schemas,
    }


def build_xlsx(results):
    wb = openpyxl.Workbook()

    # Sheet1 宽表矩阵
    ws = wb.active
    ws.title = "提取结果矩阵"
    schemas = list(SCHEMA_SLOTS.keys())
    header = ["动词", "总频次", "候选凸显图式(AI初筛)", "图式排序"] + \
             [f"S_{s[2:]}({i+1})" for i, s in enumerate(SCHEMA_SLOTS)]
    ws.append(header)
    for c in ws[1]:
        c.font = HEADER_FONT; c.fill = HEADER_FILL; c.alignment = CENTER
    for r in sorted(results, key=lambda x: -(x["freq"] or 0)):
        row = [r["verb"], r["freq"], r["candidate_schema"], " > ".join(
            s[2:] for s in r["ranking"])]
        for s in schemas:
            row.append(r["schemas"][s]["prominence_score"])
        ws.append(row)
    cand_col = 3
    for row in ws.iter_rows(min_row=2):
        row[cand_col - 1].fill = CAND_FILL
    widths = [12, 14, 22, 40] + [14] * len(SCHEMA_SLOTS)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    # Sheet2 长表明细
    ws2 = wb.create_sheet("逐词逐图式明细")
    ws2.append(["动词", "意象图式", "覆盖槽位", "覆盖率", "搭配词数N",
                "Mean_logDice", "重心得分S", "候选排名", "Top3证据(词@logDice@槽位)"])
    for c in ws2[1]:
        c.font = HEADER_FONT; c.fill = HEADER_FILL; c.alignment = CENTER
    for r in results:
        ranks = {s: i + 1 for i, s in enumerate(r["ranking"])}
        for s in SCHEMA_SLOTS:
            d = r["schemas"][s]
            top3_parts = []
            for e in d["top_evidence"][:3]:
                slot_short = e["slot"].replace(' of "%w"', "").replace('"', "")
                top3_parts.append(f"{e['word']}@{e['logdice']}@{slot_short}")
            top3 = "; ".join(top3_parts) if top3_parts else "—"
            ws2.append([r["verb"], s, f'{d["slots_present"]}/{d["slots_defined"]}',
                        d["coverage"], d["n_types"], d["mean_logdice"],
                        d["prominence_score"], ranks[s], top3])
    for i, w in enumerate([12, 16, 10, 8, 10, 13, 12, 10, 60], 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions

    # Sheet3 说明
    ws3 = wb.create_sheet("说明")
    notes = [
        ["图式提取结果数据集 —— 说明"],
        [],
        ["定位", "AI 初筛层：按《03 意象图式提取说明》2.1 的槽位映射，从 enTenTen21 Word Sketch（maxitems=50）自动聚合。未做语义归类。"],
        ["重心得分", "S(图式) = Mean_logDice × log2(N+1)；N = 该图式证据槽位去重搭配词型数（跨槽位同形词取最大 logDice）"],
        ["覆盖率", "该图式定义槽位中实际存在证据的槽位占比；覆盖率低 = 该图式证据面缺失"],
        ["候选凸显图式", "S 最大者，黄色底列；排序列给出四图式得分从高到低顺序"],
        ["⚠ 必读局限", "1) 未做语义归类：抽象/隐喻搭配词（如 reach 的 agreement/goal）未剔除，得分含多义噪声；2) 词目层数据：多义动词语料跨义位聚合；3) 槽位重合：objects/PP 同时是终点抵达与互动位移的证据面；4) 按《义项判定标准修改稿》标准2，最终图式标签须结合释义+例句人工判定，本表仅为候选参考。"],
        ["规范化路径", "人工在 03_语料搭配/<动词>.xlsx「搭配明细」I/J 列完成语义归类后，按同一公式在「保留」集合上重算，即得归类后重心得分；义位级提取按 SenseID 筛选搭配词。"],
        ["配套文件", "schema_evidence/<动词>.json：逐词机器可读结果（含每图式 Top10 证据与逐槽位明细）"],
        ["方法依据", "01_方法说明/03_意象图式提取说明.md（映射表与四步法）；04_logDice与原型性得分说明.md（公式族）"],
    ]
    for row in notes:
        ws3.append(row)
    ws3["A1"].font = Font(bold=True, size=13)
    ws3.column_dimensions["A"].width = 16
    ws3.column_dimensions["B"].width = 110
    for row in ws3.iter_rows(min_row=2):
        row[1].alignment = WRAP

    out = os.path.join(OUT_DIR, "图式提取结果汇总.xlsx")
    wb.save(out)
    return out


def main():
    os.makedirs(JSON_DIR, exist_ok=True)
    verbs = json.load(open(VERBS_FILE, encoding="utf-8"))
    results = []
    for i, v in enumerate(verbs, 1):
        r = extract_verb(v)
        results.append(r)
        with open(os.path.join(JSON_DIR, v + ".json"), "w", encoding="utf-8") as f:
            json.dump(r, f, ensure_ascii=False, indent=1)
        if i % 20 == 0:
            print(f"  {i}/{len(verbs)}")
    xlsx = build_xlsx(results)
    print("saved:", xlsx)
    print("saved per-verb JSON:", len(results), "->", JSON_DIR)
    # 概况
    from collections import Counter
    cand = Counter(r["candidate_schema"] for r in results)
    print("候选凸显图式分布:", dict(cand))


if __name__ == "__main__":
    main()
