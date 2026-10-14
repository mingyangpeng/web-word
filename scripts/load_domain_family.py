#!/usr/bin/env python3
"""把 docs/域系初判表v0.2.xlsx 灌入 MySQL 的 domain_family_senses 表。

数据链：物理位移义位表（358）→ 义位筛选规则表（R1 剔 6）→ 域系初判 v0.2（352）。
口径：词目层初判预览（义位图式暂取词目初判），最终域—系以义位层标注后重聚类为准。

用法：
    MYSQL_HOST=127.0.0.1 python3 scripts/load_domain_family.py            # 全量替换
    MYSQL_HOST=127.0.0.1 python3 scripts/load_domain_family.py --if-empty  # 仅空表时灌入
"""

import argparse
import os
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook  # noqa: E402

from src import mysql_db  # noqa: E402

XLSX = ROOT / "docs" / "域系初判表v0.2.xlsx"
WL_XLSX = ROOT / "docs" / "词表义项清单.xlsx"  # 中英对照义位页（中文释义/例句/翻译来源）


def _load_cn_index():
    """词表义项清单「中英对照义位」页 → {(word, sense_num): {cn_def, exs, cn_exs}}。

    该页来自 dictionary.db（已审校中文释义），覆盖 17 词 246 义位；
    例句与中文例句按换行拆分后逐条 zip 对齐。
    """
    if not WL_XLSX.exists():
        return {}
    wb = load_workbook(WL_XLSX, read_only=True)
    ws = wb["中英对照义位"]
    idx = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[0] or r[1] is None:
            continue
        key = (str(r[0]).strip(), int(r[1]))
        # 列序：…5 例句数, 6 例句, 7 中文例句, 8 近义…
        exs = [e for e in str(r[6] or "").split("\n") if e.strip()]
        cn_exs = [e for e in str(r[7] or "").split("\n") if e.strip()]
        idx[key] = {
            "definition_cn": str(r[2] or "").strip() or None,
            "exs": exs,
            "cn_exs": cn_exs,
        }
    wb.close()
    return idx


def read_rows() -> tuple:
    """返回 (主表行, 例句行)。"""
    cn_idx = _load_cn_index()
    wb = load_workbook(XLSX, read_only=True)
    ws = wb["02 逐义位明细"]
    rows, example_rows = [], []
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[1] is None or r[1] == "合计":
            continue
        rule = str(r[9] or "R0")
        # rule 列形态："R0" / "R2 空间位移（…）" / "边界案例：保留+需复核"（R3）
        _tag, _, _note = rule.partition(" ")
        if _tag not in ("R0", "R2", "R3"):
            _tag, _note = "R3", rule  # v0.2 表中 R3 行的规则列是描述文本
        # 02 表列序（B 起）：SenseID, 词目, 英文释义, 例句, 图式, 域, 系, 标注, 规则
        # SenseID「动词-Sn」拆分：word + sense_num 组合主键定位单词义项
        sid = str(r[1])
        word = str(r[2])
        try:
            sense_num = int(sid.rsplit("-S", 1)[1])
        except (IndexError, ValueError):
            print(f"SenseID 无法拆出序号，跳过: {sid}")
            continue

        # 例句：优先 dictionary.db 的对齐例句（含翻译），否则用 OALD 例句（无翻译）
        cn = cn_idx.get((word, sense_num))
        if cn and cn["exs"]:
            examples = list(zip(cn["exs"], cn["cn_exs"]))
            ex_source = "dictionary.db"
        else:
            oald_exs = [e for e in str(r[4] or "").split("\n") if e.strip()]
            examples = [(e, None) for e in oald_exs]
            ex_source = "OALD10"
        for i, (en, zh) in enumerate(examples, start=1):
            example_rows.append({
                "word": word, "sense_num": sense_num, "ex_num": i,
                "example_en": en, "example_cn": zh, "source": ex_source,
            })
        first_en = examples[0][0] if examples else None

        rows.append({
            "word": word,
            "sense_num": sense_num,
            "definition": str(r[3] or "") or None,
            "definition_cn": (cn or {}).get("definition_cn"),
            "example": first_en,
            "lemma_schema": str(r[5] or "") or None,
            "domain": str(r[6]),
            "family": str(r[7]),
            "sense_annotation": str(r[8] or "") or None,
            "rule_tag": _tag,
            "rule_note": _note or None,
            "source": "域系初判表v0.2",
            "loaded_at": date.today(),
        })
    wb.close()
    return rows, example_rows


def main() -> int:
    parser = argparse.ArgumentParser(description="灌入域—系初判表 v0.2")
    parser.add_argument("--if-empty", action="store_true", help="仅当表为空时写入")
    args = parser.parse_args()

    if not XLSX.exists():
        print(f"源文件不存在: {XLSX}")
        return 1

    if args.if_empty and mysql_db.domain_family_ready():
        print("domain_family_senses 已有数据，跳过（--if-empty）。")
        return 0

    rows, example_rows = read_rows()
    if not rows:
        print("未读到任何行，中止。")
        return 1

    ok = mysql_db.init_database()
    if not ok:
        print("数据库初始化失败。")
        return 1

    n = mysql_db.replace_domain_family_senses(rows, example_rows)
    n_ex = len(example_rows)
    n_cn = sum(1 for r in rows if r.get("definition_cn"))
    print(f"已灌入 domain_family_senses: {n} 行（来源 {XLSX.name}）")
    print(f"  例句子表 domain_family_examples: {n_ex} 行")
    print(f"  含中文释义的义位: {n_cn}/{n}（dictionary.db 覆盖 17 词）")
    return 0 if n == len(rows) else 1


if __name__ == "__main__":
    sys.exit(main())
