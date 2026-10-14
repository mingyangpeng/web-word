#!/usr/bin/env python3
"""导出「域—系初判表（v0.1 词目层初判版）」——358 候选义位按图式聚系归域的预览。

⚠️ 口径声明（必须随表传播）：
  本表是【词目层初判】的预览：义位的图式归属暂取其词目在筛选表中的初判图式
  （SCHEMA_BY_SUB_CAT 映射），不是义位层标注结果。真正的域—系体系（论文 4.2.2）
  必须在 LLM 初筛 + 人工逐义位复核（编码手册流程）完成后，由 Y 类义位的实际
  图式标签重新聚类生成。本表用途：预览体系骨架、暴露词目层初判的错位
  （如 bolt-S1 闩门义混进奔跑类），为标注提供对照基线。

产出：docs/域系初判表.xlsx
  01 域—系总览：域 → 系 → 义位数 → 代表例句
  02 逐义位明细：SenseID、释义、例句、词目层初判图式、域/系归属
  03 口径与错位说明
"""

import os
import sys
from datetime import date
from pathlib import Path
from collections import defaultdict

XLSX_SKILL_DIR = "/Users/jhon/.zcode/cli/plugins/cache/zcode-plugins-official/spreadsheets/0.1.7/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)
from base import (  # noqa: E402
    FONT_NAME, FORMATS, PRIMARY_LIGHT, setup_sheet, style_header_row,
    style_data_row, style_total_row, font_caption, font_subheader,
    auto_fit_columns, auto_fit_row_heights,
)
from openpyxl import Workbook, load_workbook  # noqa: E402
from openpyxl.styles import Font, PatternFill  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "docs" / "域系初判表.xlsx"

# ---- 数据 ----
phys = load_workbook(ROOT / "docs" / "物理位移义位表.xlsx", read_only=True)
cand = []
for r in phys["物理位移义位"].iter_rows(min_row=5, values_only=True):
    if r[1] is None or r[1] == "合计":
        break
    cand.append({"verb": str(r[1]), "sid": f"{r[1]}-S{r[2]}",
                 "defn": str(r[6] or ""), "ex": str(r[8] or "")})
phys.close()

lex = {}
wl = load_workbook(ROOT / "docs" / "词表义项清单.xlsx", read_only=True)
for r in wl["词表概览"].iter_rows(min_row=2, values_only=True):
    if r[0]:
        lex[str(r[0]).strip()] = (r[2], r[3], r[4])
wl.close()

DOMAIN = {"方式伴随图式": "方式位移域", "终点抵达图式": "路径位移域",
          "源点脱离图式": "路径位移域", "互动位移图式": "互动位移域"}
for c in cand:
    cat, sub, schema = lex.get(c["verb"], ("?", "?", "?"))
    c.update(sub=sub, schema=schema, domain=DOMAIN.get(schema, "?"))

by_dom = defaultdict(lambda: defaultdict(list))
for c in cand:
    by_dom[c["domain"]][c["sub"]].append(c)

# ---- 工作簿 ----
wb = Workbook()

# 01 域—系总览
ws = wb.active
ws.title = "01 域系总览"
H1 = ["认知域（初判）", "语义系（初判）", "义位数", "代表义位与例句"]
setup_sheet(ws, title="域—系初判表 v0.1（词目层初判 · 358 候选义位 · 非最终体系）", last_col=len(H1) + 1)
for col, h in enumerate(H1, 2):
    ws.cell(row=4, column=col, value=h)
style_header_row(ws, row_num=4, col_start=2, col_end=len(H1) + 1)
rn = 5
i = 0
for dom in ["方式位移域", "路径位移域", "互动位移域"]:
    subs = by_dom.get(dom, {})
    for sub, items in sorted(subs.items(), key=lambda x: -len(x[1])):
        exs = []
        for it in items:
            if it["ex"] and it["ex"] != "None" and len(exs) < 2:
                first = it["ex"].split("\n")[0]
                exs.append(f"{it['sid']}: {first}")
        ws.cell(row=rn, column=2, value=dom)
        ws.cell(row=rn, column=3, value=sub)
        ws.cell(row=rn, column=4, value=len(items))
        ws.cell(row=rn, column=5, value="\n".join(exs))
        style_data_row(ws, row_num=rn, col_start=2, col_end=len(H1) + 1, row_index=i)
        ws.cell(row=rn, column=4).number_format = FORMATS["integer"]
        rn += 1
        i += 1
    # 域小计
    dom_total = sum(len(v) for v in subs.values())
    ws.cell(row=rn, column=2, value=f"{dom} 小计")
    ws.cell(row=rn, column=4, value=dom_total)
    style_total_row(ws, row_num=rn, col_start=2, col_end=len(H1) + 1)
    ws.cell(row=rn, column=4).number_format = FORMATS["integer"]
    rn += 1
ws.cell(row=rn, column=2, value="合计")
ws.cell(row=rn, column=4, value=len(cand))
style_total_row(ws, row_num=rn, col_start=2, col_end=len(H1) + 1)
ws.cell(row=rn, column=4).number_format = FORMATS["integer"]
ws.cell(row=rn + 2, column=2,
        value="注：本表为词目层初判预览——义位图式暂取词目初判，最终体系以义位层人工标注后重聚类为准（见 03 页）。"
              "已知错位示例：bolt-S1（闩门）混入奔跑类、bike-S2（邮寄义）混入运载工具类，正是义位层复核要修正的对象。").font = font_caption()
ws.column_dimensions["E"].width = 80
auto_fit_columns(ws, min_width=8, max_width=30, header_row=4, data_start_row=5)
auto_fit_row_heights(ws, header_row=4, data_start_row=5)

# 02 逐义位明细
ws2 = wb.create_sheet("02 逐义位明细")
H2 = ["SenseID", "词目", "英文释义", "例句（首条）", "词目层初判图式", "域（初判）", "系（初判）", "义位层标注（待人工）"]
setup_sheet(ws2, title="358 候选义位逐条明细（含域—系初判归属）", last_col=len(H2) + 1)
for col, h in enumerate(H2, 2):
    ws2.cell(row=4, column=col, value=h)
style_header_row(ws2, row_num=4, col_start=2, col_end=len(H2) + 1)
for i, c in enumerate(sorted(cand, key=lambda x: ({"方式位移域":0,"路径位移域":1,"互动位移域":2}[x["domain"]], x["sub"], x["verb"]))):
    rn = 5 + i
    first_ex = (c["ex"].split("\n")[0] if c["ex"] and c["ex"] != "None" else "")
    vals = [c["sid"], c["verb"], c["defn"], first_ex, c["schema"], c["domain"], c["sub"], None]
    for col, v in enumerate(vals, 2):
        ws2.cell(row=rn, column=col, value=v)
    style_data_row(ws2, row_num=rn, col_start=2, col_end=len(H2) + 1, row_index=i)
    ws2.cell(row=rn, column=9).fill = PatternFill("solid", fgColor=PRIMARY_LIGHT)
total = 5 + len(cand)
ws2.cell(row=total, column=2, value="合计")
ws2.cell(row=total, column=3, value=f"{len(cand)} 义位")
style_total_row(ws2, row_num=total, col_start=2, col_end=len(H2) + 1)
ws2.freeze_panes = "D5"
ws2.auto_filter.ref = f"B4:I{total - 1}"
auto_fit_columns(ws2, min_width=8, max_width=56, header_row=4, data_start_row=5)
auto_fit_row_heights(ws2, header_row=4, data_start_row=5)

# 03 口径与错位说明
ws3 = wb.create_sheet("03 口径与错位说明")
setup_sheet(ws3, title="域—系初判表 · 口径声明与已知错位", last_col=3)
ws3.column_dimensions["B"].width = 24
ws3.column_dimensions["C"].width = 100
rn = 4

def kv(k, v, warn=False):
    global rn
    ws3.cell(row=rn, column=2, value=k)
    ws3.cell(row=rn, column=3, value=v)
    style_data_row(ws3, row_num=rn, col_start=2, col_end=3, row_index=rn % 2)
    if warn:
        ws3.cell(row=rn, column=3).font = Font(name=FONT_NAME, size=11, color="C0392B")
    rn += 1

kv("本表性质", "词目层初判预览（v0.1）：义位的图式归属暂取其词目的筛选表初判图式，"
            "不是义位层标注结果，不能直接作为论文 4.2.2 的最终体系。", warn=True)
kv("生成链", "物理位移义位表（358 候选，机械正则）→ 按词目初判图式分组 → 域归并"
           "（方式→方式位移域；终点+源点→路径位移域；互动→互动位移域）。")
kv("域归并依据", "四图式→三域：方式伴随独立成域；终点抵达与源点脱离同属路径位移"
              "（事件结构互补：指向终点 vs 脱离源点）；互动位移独立成域。此为初判归并，"
              "最终域的划分以义位层归纳为准（若数据支持，源点/终点亦可各自成域）。")
kv("已知错位（词目层初判的局限证据）", "bolt-S1/S2（闩门、栓接）混入奔跑类——bolt 的位移义（飞奔）"
                              "只是多义之一；bike-S2（邮寄义）混入运载工具类；advance-S7（晋级义）"
                              "混入终点类。这些错位正是论文 4.2.3 的素材：词目层初判无法表达义位层异质性。")
kv("下一步（正式体系生成）", "①LLM 初筛 358 条 → ②人工逐义位复核（Y/拓展/N＋四图式）→ "
                     "③Y 类义位按实际图式标签重新聚类成系、归并成域 → ④重跑本脚本升级为 v1.0 正式体系表。")
kv("生成信息", f"生成日期 {date.today()}；脚本 scripts/export_domain_family_table.py 可重跑。")
auto_fit_row_heights(ws3, header_row=4, data_start_row=4)

wb.properties.creator = "Z.ai"
wb.save(OUT)
print(f"saved: {OUT}")
print(f"total={len(cand)}; domains: 方式{sum(len(v) for s in by_dom['方式位移域'].values() for v in [s])}"
      f"/路径{sum(len(v) for s in by_dom['路径位移域'].values() for v in [s])}"
      f"/互动{sum(len(v) for s in by_dom['互动位移域'].values() for v in [s])}")
