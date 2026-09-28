"""
导出词表义项清单为 xlsx，供人工核查。

用户要求「同时输出这些词的义项到 xlsx，方便我检查」。这个脚本把灌库前的**全部**
义项摊平成表 —— 不是为了好看，而是为了让核查能落到具体某一行的某一列上：
哪条释义被清洗坏了、哪条习语义项该剔除、哪个词一条义项都没有，都要能一眼定位。

四张表各有分工：

| 工作表 | 一行是什么 | 用来核查什么 |
|--------|-----------|--------------|
| `词表概览` | 一个词 | 覆盖是否完整：哪些词一条义项都没有、来源分布、音标缺口 |
| `义项明细` | 一个义项 | **主表**。清洗对不对、语法码有没有剥干净、习语要不要剔除 |
| `中英对照义位` | 一个 dictionary.db 义位 | 中文释义是否可用作词目释义（与上表是两套切分体系，故分开） |
| `问题清单` | 一个待确认点 | 脚本自己发现的可疑之处，不用人去猜哪里可能有问题 |

**所有列都是原始值加最小加工**：例句在 `义项明细` 里合并成一列用换行分隔，同时另给
`例句数` 列 —— Excel 里换行显示会挤，但拆成多列会让某些义项（`run` 有 8 句）撑出
很宽的稀疏表。核查时按 `例句数` 升序排一遍就能先看例句少的那些。
"""

import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)

from src.oald_senses import (  # noqa: E402
    annotate_oald_senses,
    audit_coverage as oald_audit,
    load_oald_verb_index,
    senses_for as oald_senses_for,
)
from src.word_source import (  # noqa: E402
    DictionaryDB,
    build_entries,
    load_legacy_library,
    load_levin_table,
    load_oxford_cache,
)

DEFAULT_OUT = os.path.join(ROOT, "docs/词表义项清单.xlsx")


def collect():
    """把三个来源读成「词 -> 全部义项」的结构，不做任何过滤。"""
    levin = load_levin_table()          # 默认只取「方式聚焦型」136 词
    db = DictionaryDB()
    index = load_oald_verb_index()

    # 跑一遍真实的释义适配，拿到灌库时**实际会写入**的中文释义与来源。
    # 不在导出里另算一遍 —— 另算就多了一套会漂移的逻辑，而这一页的用途正是核对它。
    built = {e["verb"]: e for e in build_entries(
        levin, load_oxford_cache(), db, load_legacy_library()
    )}

    rows = []
    for entry in levin:
        verb = entry["verb"]
        dict_senses = db.senses(verb) if db.available else []
        oald, how, names = oald_senses_for(verb, index)
        # 标注与灌库走同一个函数。导出里另算一遍就会漂移，而这一页正是用来核对它的。
        entry["_has_dict_senses"] = bool(dict_senses)
        oald = annotate_oald_senses(oald, entry)
        b = built.get(verb, {})
        rows.append({
            "levin": entry,
            "verb": verb,
            "phonetic": db.phonetic(verb) if db.available else "",
            "chinese_ndef": b.get("chinese_ndef") or "",
            "definition_source": b.get("definition_source") or "",
            "dict_senses": dict_senses,
            "oald_senses": oald,
            "resolution": how,
            "headwords": names,
        })
    return rows, index


def write_workbook(rows, index, out_path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="4E79A7")
    wrap = Alignment(vertical="top", wrap_text=True)

    def style_header(ws, widths):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for cell in ws[1]:
            cell.font = head_font
            cell.fill = head_fill
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.freeze_panes = "A2"

    # ---------- 1. 词表概览 ----------
    ws = wb.active
    ws.title = "词表概览"
    ws.append(["词目", "音标", "主范畴", "次类", "意象图式", "原型主导义简述",
               "词目中文释义（灌库值）", "词目释义来源",
               "dictionary.db 义位数", "OALD 义项数", "义项合计",
               "OALD 匹配方式", "OALD 来源词条"])
    for r in rows:
        e = r["levin"]
        ws.append([
            r["verb"], r["phonetic"], e.get("main_cat"), e.get("sub_cat"),
            e.get("image_schema"), e.get("sense_brief"),
            r["chinese_ndef"], r["definition_source"],
            len(r["dict_senses"]), len(r["oald_senses"]),
            len(r["dict_senses"]) + len(r["oald_senses"]),
            {"direct": "直接命中", "alias": "同形异义归并", "homograph": "编号词条",
             "missing": "表内无此词"}.get(r["resolution"], r["resolution"]),
            ", ".join(r["headwords"]),
        ])
    style_header(ws, [12, 20, 12, 14, 14, 22, 46, 14, 18, 12, 10, 16, 14])

    # ---------- 2. 义项明细 ----------
    ws2 = wb.create_sheet("义项明细")
    ws2.append(["词目", "义项序号", "来源", "语义分组", "语法码", "英文释义",
                "例句数", "例句", "句法模式", "是否习语", "原型义", "来源词条"])
    for r in rows:
        for i, s in enumerate(r["oald_senses"], start=1):
            ws2.append([
                r["verb"], i, s.get("source"), s.get("sense_group"),
                s.get("grammar"), s.get("en_ndef"),
                len(s.get("examples") or []),
                "\n".join(s.get("examples") or []),
                # 只写有模式的那几行：没有模式时 OALD 是「继承上一句」，
                # 直接 join 会留下一串空行，看起来像数据坏了
                "\n".join(p for p in (s.get("example_patterns") or []) if p),
                "是" if s.get("is_idiom") else "",
                "是" if s.get("prototype_level") == "Y" else "",
                s.get("headword"),
            ])
    style_header(ws2, [12, 10, 12, 18, 26, 64, 8, 70, 22, 10, 8, 12])
    for row in ws2.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap

    # ---------- 3. 中英对照义位（dictionary.db）----------
    ws3 = wb.create_sheet("中英对照义位")
    ws3.append(["词目", "义位序号", "中文释义", "英文释义", "句式", "例句数",
                "例句", "中文例句", "近义", "反义"])
    for r in rows:
        for i, s in enumerate(r["dict_senses"], start=1):
            ws3.append([
                r["verb"], i, s.get("ch_ndef"), s.get("en_ndef"),
                s.get("construction"), len(s.get("examples") or []),
                "\n".join(s.get("examples") or []),
                "\n".join(s.get("ch_examples") or []),
                ", ".join(s.get("synonyms") or []),
                ", ".join(s.get("antonyms") or []),
            ])
    style_header(ws3, [12, 10, 46, 46, 20, 8, 62, 40, 26, 20])
    for row in ws3.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap

    # ---------- 4. 问题清单 ----------
    #
    # 这一页是脚本**自己发现的可疑点**，不是等人工去找。宁可多列几条待确认，
    # 也不要因为「看起来没问题」而漏掉一整类词 —— 实测 missing 那 4 个词
    # 就是靠这个机制冒出来的。
    ws4 = wb.create_sheet("问题清单")
    ws4.append(["类型", "词目", "说明"])
    for r in rows:
        if r["resolution"] == "missing":
            ws4.append(["表内无此词", r["verb"],
                        "OALD 导出表中查不到该词的动词义项（已逐个复核：或只有名词词条，"
                        "或词条名都不存在）。该词在义项明细里为 0 行。"])
        elif r["resolution"] in ("alias", "homograph"):
            ws4.append(["同形异义归并", r["verb"],
                        f"该词在表中是带编号的同形异义词条 {r['headwords']}，"
                        f"义项已合并到本词目下，可用「来源词条」列区分。"])
        for i, s in enumerate(r["oald_senses"], start=1):
            if s.get("is_idiom"):
                ws4.append(["习语义项", r["verb"],
                            f"第 {i} 条义项属于 IDIOMS 组，释义是整句习语的解释而非该词的"
                            f"位移义（{s.get('definition', '')[:60]}…）。默认全部保留，"
                            f"是否剔除请你在这一页确认。"])
            if not s.get("examples"):
                ws4.append(["无例句", r["verb"],
                            f"第 {i} 条义项没有例句：{s.get('definition', '')[:60]}…"])
        if not r["phonetic"]:
            ws4.append(["无音标", r["verb"], "dictionary.db 未收录该词音标。"])
        if not r["dict_senses"] and not r["oald_senses"]:
            ws4.append(["零义项", r["verb"], "两个来源都没有该词的任何义项。"])
    style_header(ws4, [16, 12, 110])
    for row in ws4.iter_rows(min_row=2):
        for c in row:
            c.alignment = wrap

    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    wb.save(out_path)
    return out_path


def main():
    import argparse
    parser = argparse.ArgumentParser(description="导出词表义项清单 xlsx（供人工核查）")
    parser.add_argument("-o", "--output", default=DEFAULT_OUT, help="输出路径")
    args = parser.parse_args()

    rows, index = collect()
    verbs = [r["verb"] for r in rows]
    report = oald_audit(verbs, index)

    path = write_workbook(rows, index, args.output)
    n_oald = sum(len(r["oald_senses"]) for r in rows)
    n_dict = sum(len(r["dict_senses"]) for r in rows)

    print(f"已导出: {os.path.relpath(path, ROOT)}")
    print(f"  词目 {len(rows)} 个")
    print(f"  OALD 义项 {n_oald} 条（习语 {report['idiom_senses']}、"
          f"带例句 {report['senses_with_examples']}、带语法码 {report['senses_with_grammar']}）")
    print(f"  dictionary.db 义位 {n_dict} 条")
    print(f"  解析方式 {report['by_resolution']}")
    if report["missing_verbs"]:
        print(f"  ⚠️  表内无动词义项: {report['missing_verbs']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
