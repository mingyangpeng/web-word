#!/usr/bin/env python3
"""导出「词表 12 次类构成表」——三主范畴 × 十二次类的词池构成总表。

数据来源：
  - docs/Levin移动动词候选词筛选表.xlsx「去除不符合的」页（主范畴/次类/词目，权威源）
  - output/02_义位标注/*_义项.xlsx（每词义位数，合计 996）
  - src/word_source.py SCHEMA_BY_SUB_CAT（词目层意象图式初判映射）

产出：docs/词表12次类构成表.xlsx，供论文 4.1.1（词池构成）直接引用。
"""

import glob
import os
import sys
from collections import defaultdict
from datetime import date
from pathlib import Path

XLSX_SKILL_DIR = "/Users/jhon/.zcode/cli/plugins/cache/zcode-plugins-official/spreadsheets/0.1.7/skills/xlsx"
for sub in [XLSX_SKILL_DIR, os.path.join(XLSX_SKILL_DIR, "templates")]:
    if sub not in sys.path:
        sys.path.insert(0, sub)
from base import (  # noqa: E402
    FORMATS, setup_sheet, style_header_row, style_data_row, style_total_row,
    font_caption, auto_fit_columns, auto_fit_row_heights,
)
from openpyxl import Workbook, load_workbook  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
OUT = ROOT / "docs" / "词表12次类构成表.xlsx"

wb = load_workbook(ROOT / "docs" / "Levin移动动词候选词筛选表.xlsx", read_only=True)
ws = wb["去除不符合的"]
h = {c.value: i for i, c in enumerate(ws[1])}
groups = defaultdict(list)
for r in ws.iter_rows(min_row=2, values_only=True):
    if not r[0]:
        continue
    groups[(str(r[h["主范畴"]]), str(r[h["次类"]]))].append(str(r[h["词目(lemma)"]]).strip())
wb.close()

sense = {}
for f in glob.glob(str(ROOT / "output" / "02_义位标注" / "*_义项.xlsx")):
    w = load_workbook(f, read_only=True)
    n = sum(1 for r in w["义位标注"].iter_rows(min_row=2, values_only=True) if r[0])
    sense[os.path.basename(f).replace("_义项.xlsx", "")] = n
    w.close()

from src.word_source import SCHEMA_BY_SUB_CAT  # noqa: E402  图式映射单一来源

rows = []
for (cat, sub), words in groups.items():
    rows.append((cat, sub, len(words), sum(sense.get(w, 0) for w in words),
                 SCHEMA_BY_SUB_CAT[sub], "、".join(sorted(words))))
rows.sort(key=lambda x: ({"方式聚焦型": 0, "路径聚焦型": 1, "互动聚焦型": 2}[x[0]], -x[2]))

wb2 = Workbook()
ws2 = wb2.active
ws2.title = "次类构成"
HEADERS = ["主范畴", "次类", "词数", "义位数", "词目层意象图式初判", "全部词目（按字母序）"]
LAST = len(HEADERS) + 1
setup_sheet(ws2, title="词表 12 次类构成表（三主范畴 × 十二次类 · 171 词 / 996 义位）", last_col=LAST)
for col, hd in enumerate(HEADERS, 2):
    ws2.cell(row=4, column=col, value=hd)
style_header_row(ws2, row_num=4, col_start=2, col_end=LAST)
for i, r in enumerate(rows):
    rn = 5 + i
    for col, v in enumerate(r, 2):
        ws2.cell(row=rn, column=col, value=v)
    style_data_row(ws2, row_num=rn, col_start=2, col_end=LAST, row_index=i)
    for col in (4, 5):
        ws2.cell(row=rn, column=col).number_format = FORMATS["integer"]
total = 5 + len(rows)
ws2.cell(row=total, column=2, value="合计")
ws2.cell(row=total, column=4, value=sum(r[2] for r in rows))
ws2.cell(row=total, column=5, value=sum(r[3] for r in rows))
style_total_row(ws2, row_num=total, col_start=2, col_end=LAST)
for col in (4, 5):
    ws2.cell(row=total, column=col).number_format = FORMATS["integer"]
ws2.cell(row=total + 2, column=2,
         value=f"来源：Levin 移动动词候选词筛选表「去除不符合的」页（171 词 = Talmy 语义核查与构式交替核查通过者）；"
               f"义位数取自 02_义位标注（992 条 OALD10 义项 + 4 条补充来源）；词目层意象图式初判为规则映射候选，"
               f"义位层归属以标注为准。生成日期 {date.today()}；脚本 scripts/export_subcategory_table.py 可重跑。"
               f"本表可直接作为论文 4.1.1 的词池构成表。").font = font_caption()
ws2.freeze_panes = "C5"
auto_fit_columns(ws2, min_width=8, max_width=90, header_row=4, data_start_row=5)
auto_fit_row_heights(ws2, header_row=4, data_start_row=5)

wb2.properties.creator = "Z.ai"
wb2.save(OUT)
print(f"saved: {OUT}")
print(f"rows={len(rows)}, words={sum(r[2] for r in rows)}, senses={sum(r[3] for r in rows)}")
